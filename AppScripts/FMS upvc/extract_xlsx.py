# -*- coding: utf-8 -*-
"""Extract master worksheets from the source workbook into new database schema CSV text files.

Built strictly according to the uPVC Database Architecture plan:
- Section 1: uPVC_Quotation_Master (15 columns)
- Section 3: uPVC_Purchase_Order_Master (19 columns)

Usage:
    python extract_xlsx.py
"""
import csv
import datetime
import os
import re

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, 'uPVC PO _ Quotation FMS.xlsx')

QUOTATION_HEADERS = [
    'Quotation ID',
    'Client / Site Name',
    'Date',
    'Location',
    'Vendor Name',
    'Variant',
    'Mesh / Jaali',
    'Glass Specs',
    'Total Item Qty',
    'Lead Time',
    'Rough Sizes',
    'Vendor Quote',
    'Vendor Amount',
    'Our Quote',
    'Our Amount',
]

PO_HEADERS = [
    'PO ID',
    'Quotation ID',
    'Client / Site Name',
    'PO Date',
    'Location',
    'Vendor Name',
    'Variant',
    'Mesh / Jaali',
    'Glass Specs',
    'Total Item Qty',
    'Lead Time (Days)',
    'Rough Sizes',
    'Vendor Quote',
    'Vendor Amount',
    'Our Quote',
    'Our Amount',
    'PO Order',
    'Bills',
    'Bills AMT',
]


def format_date(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    m = re.match(r'^(\d{4}-\d{2}-\d{2})', s)
    if m:
        return m.group(1)
    return s


def format_doc_url(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s.startswith('http://') or s.startswith('https://'):
        return s
    return ''


def format_num(v):
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    s = str(v).strip()
    # Check if numeric
    clean = re.sub(r'[^0-9.]', '', s)
    return clean if clean else ''


def format_int(v):
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        return str(int(round(v)))
    nums = re.findall(r'\d+', str(v))
    if nums:
        return str(sum(int(n) for n in nums))
    return ''


def build_mesh(with_mesh_val, mesh_type_val):
    w = str(with_mesh_val or '').strip().lower()
    t = str(mesh_type_val or '').strip()
    if t:
        return t if 'mesh' in t.lower() else f'{t} Mesh'
    if w in ('no', 'none'):
        return 'None'
    if w in ('yes', 'y'):
        return 'Yes'
    return 'None' if not w else w.capitalize()


def build_glass_specs(thickness_val, color_val, toughened_val):
    thick = str(thickness_val or '').strip()
    color = str(color_val or '').strip()
    tough = str(toughened_val or '').strip().lower()
    parts = []
    if thick:
        parts.append(thick)
    if color:
        parts.append(color.capitalize() if color.islower() else color)
    if tough in ('yes', 'y', 'toughened'):
        parts.append('Toughened')
    return ', '.join(parts)


def extract_quotations(wb):
    ws = wb['uPVC Quotation']
    data_rows = []
    qt_site_to_id = {}

    rows = list(ws.iter_rows(values_only=True))
    # Row 6 (0-indexed 5) is header row, row 7+ is data
    data_start = 6
    idx = 1
    for r in rows[data_start:]:
        site = r[0]
        if not site or not str(site).strip():
            continue

        site_str = str(site).strip()
        qt_id = f'QT-UPVC-{idx:04d}'
        qt_site_to_id[re.sub(r'\s+', ' ', site_str.lower())] = qt_id
        idx += 1

        date_str = format_date(r[1])
        location = str(r[2]).strip() if r[2] else ''
        vendor = str(r[3]).strip() if r[3] else ''
        variant = str(r[5]).strip() if r[5] else ''

        mesh = build_mesh(r[7], r[11])
        glass = build_glass_specs(r[8], r[12], r[9])
        total_qty = ''
        lead_time = format_int(r[10])

        rough_sizes = format_doc_url(r[4])
        vendor_quote = format_doc_url(r[13])
        vendor_amount = ''
        our_quote = format_doc_url(r[14])
        our_amount = ''

        data_rows.append([
            qt_id,
            site_str,
            date_str,
            location,
            vendor,
            variant,
            mesh,
            glass,
            total_qty,
            lead_time,
            rough_sizes,
            vendor_quote,
            vendor_amount,
            our_quote,
            our_amount,
        ])

    return data_rows, qt_site_to_id


def extract_purchase_orders(wb, qt_site_to_id):
    ws = wb['uPVC PO']
    data_rows = []

    rows = list(ws.iter_rows(values_only=True))
    data_start = 6
    idx = 1
    for r in rows[data_start:]:
        site = r[0]
        if not site or not str(site).strip():
            continue

        site_str = str(site).strip()
        po_id = f'PO-UPVC-{idx:04d}'
        idx += 1

        norm_site = re.sub(r'\s+', ' ', site_str.lower())
        qt_id = qt_site_to_id.get(norm_site, '')

        po_date = format_date(r[1])
        vendor = str(r[2]).strip() if r[2] else ''
        location = str(r[3]).strip() if r[3] else ''

        # Variant + Shutter
        variant_parts = []
        if r[6] and str(r[6]).strip():
            variant_parts.append(str(r[6]).strip())
        if r[8] and str(r[8]).strip() and str(r[8]).strip().lower() not in (str(r[6] or '').lower()):
            variant_parts.append(str(r[8]).strip())
        variant = ', '.join(variant_parts)

        mesh = build_mesh(r[15], r[16])
        glass = build_glass_specs(r[9], r[10], r[11])

        # Window Qty + Ventilator Qty
        w_qty = 0
        v_qty = 0
        if r[17] is not None:
            if isinstance(r[17], (int, float)):
                w_qty = int(round(r[17]))
            else:
                nums = re.findall(r'\d+', str(r[17]))
                w_qty = sum(int(n) for n in nums)
        if r[18] is not None:
            if isinstance(r[18], (int, float)):
                v_qty = int(round(r[18]))
            else:
                nums = re.findall(r'\d+', str(r[18]))
                v_qty = sum(int(n) for n in nums)
        total_qty = str(w_qty + v_qty) if (w_qty or v_qty) else ''

        lead_time = format_int(r[12])
        rough_sizes = format_doc_url(r[4])
        vendor_quote = format_doc_url(r[5])
        vendor_amount = format_num(r[13])
        our_quote = format_doc_url(r[21])
        our_amount = format_num(r[14])
        po_order = ''
        bills = format_doc_url(r[20])
        bills_amt = format_num(r[19])

        data_rows.append([
            po_id,
            qt_id,
            site_str,
            po_date,
            location,
            vendor,
            variant,
            mesh,
            glass,
            total_qty,
            lead_time,
            rough_sizes,
            vendor_quote,
            vendor_amount,
            our_quote,
            our_amount,
            po_order,
            bills,
            bills_amt,
        ])

    return data_rows


def write_csv(filename, headers, rows):
    out_path = os.path.join(HERE, filename)
    with open(out_path, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh, quoting=csv.QUOTE_MINIMAL, lineterminator='\n')
        w.writerow(headers)
        for row in rows:
            w.writerow(row)
    print(f'Wrote {filename}: {len(rows)} records, {len(headers)} columns')


def main():
    print(f'Loading workbook: {XLSX}')
    wb = load_workbook(XLSX, data_only=True, read_only=True)

    qt_rows, qt_site_to_id = extract_quotations(wb)
    write_csv('uPVC_Quotation_Master.txt', QUOTATION_HEADERS, qt_rows)

    po_rows = extract_purchase_orders(wb, qt_site_to_id)
    write_csv('uPVC_Purchase_Order_Master.txt', PO_HEADERS, po_rows)

    wb.close()
    print('Extraction complete!')


if __name__ == '__main__':
    main()
